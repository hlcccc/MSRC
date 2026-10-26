"""Maintainer-only Markdown-to-Word builder (requires python-docx)."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


def build(source, target, template=None):
    document = Document()
    formats = {}
    if template is not None:
        original = Document(template)
        document.styles.element.clear()
        for element in original.styles.element:
            document.styles.element.append(deepcopy(element))
        section = deepcopy(original.element.body.sectPr)
        for child in list(section):
            if child.tag in {qn('w:headerReference'), qn('w:footerReference'), qn('w:pgNumType')}:
                section.remove(child)
        document.element.body.replace(document.element.body.sectPr, section)
        paragraphs = original.paragraphs
        formats = dict(title=paragraphs[0]._p, body=paragraphs[1]._p,
            step=next(p._p for p in paragraphs if p.text.startswith('（1）')),
            command=next(p._p for p in paragraphs if p.text.startswith('git clone')),
            success=next(p._p for p in paragraphs if p.text.startswith('成功信号：')))

    def add(text, kind):
        paragraph = document.add_paragraph()
        if kind in formats:
            model = formats[kind]
            properties = model.find(qn('w:pPr'))
            if properties is not None:
                paragraph._p.append(deepcopy(properties))
            model_run = model.find(qn('w:r'))
            run_properties = model_run.find(qn('w:rPr')) if model_run is not None else None
        else:
            run_properties = None
            if kind in ('title', 'step'):
                paragraph.style = 'Title' if kind == 'title' else 'Heading 1'
        for piece in re.split(r'(`[^`]+`)', text) if kind != 'command' else [text]:
            if not piece:
                continue
            run = paragraph.add_run(piece.strip('`') if piece.startswith('`') else piece)
            if run_properties is not None:
                run._r.insert(0, deepcopy(run_properties))
            if kind == 'command' or piece.startswith('`'):
                run.font.name = 'Consolas'

    text = source.read_text(encoding='utf-8')
    pending, code, in_code = [], [], False
    def flush():
        if pending:
            value = ' '.join(pending)
            value = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', value)
            add(value, 'success' if value.startswith(('成功信号：', '完成信号：')) else 'body')
            pending.clear()
    for line in text.splitlines():
        if line.startswith('```'):
            flush()
            if in_code:
                add('\n'.join(code), 'command')
                code.clear()
            in_code = not in_code
        elif in_code:
            code.append(line)
        elif line.startswith('# '):
            flush()
            add(line[2:], 'title')
        elif line.startswith('## '):
            flush()
            add(line[3:], 'step')
        elif not line.strip():
            flush()
        else:
            pending.append(line)
    flush()
    assert not in_code
    document.core_properties.title = text.splitlines()[0][2:]
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.save(target)
    reread = Document(target)
    rendered = '\n'.join(p.text for p in reread.paragraphs)
    commands = re.findall(r'```bash\n(.*?)\n```', text, flags=re.S)
    assert len(commands) == 6 and all(command in rendered for command in commands)
    assert sum(p.text.startswith('（') for p in reread.paragraphs) == 6
    return dict(six_steps=True, commands_match_source=True, original_template_formats_used=bool(template),
                source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                docx_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                qa_scope='OOXML structure and exact command text; no native pagination rendering')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    source = root / 'docs/FLOW_SAFETY_MULTILAYER.md'
    result = build(source, source.with_suffix('.docx'), args.template)
    target = root / 'resources/multilayer/document_qa.json'
    target.write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result))
